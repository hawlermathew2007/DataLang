"""Read unstructured documents (PDF / DOCX / TXT / MD) into plain-text blocks.

Local only. Scanned PDFs (no text layer) are reported as such: OCR is not included.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


@dataclass
class Block:
    text: str
    page: int | None = None     # 1-based page for PDFs, None otherwise
    index: int = 0              # order in the document


class DocumentError(Exception):
    pass


def extract_blocks(name: str, data: bytes | str | Path) -> list[Block]:
    """Return paragraph-level blocks. Raises DocumentError with a readable message."""
    if isinstance(data, (str, Path)):
        data = Path(data).read_bytes()
    lower = name.lower()
    if lower.endswith(".pdf"):
        raw = _pdf(data)
    elif lower.endswith(".docx"):
        raw = _docx(data)
    elif lower.endswith((".txt", ".md")):
        raw = [(None, _decode(data))]
    else:
        raise DocumentError(f"Unsupported file type: {Path(name).suffix or name}. Supported: {', '.join(SUPPORTED)}")
    blocks: list[Block] = []
    for page, text in raw:
        for para in re.split(r"\n\s*\n", text):
            para = re.sub(r"[ \t]+", " ", para).strip()
            if para:
                blocks.append(Block(para, page, len(blocks)))
    if not blocks:
        raise DocumentError("No text found. If this is a scanned PDF/image it needs OCR first (not included).")
    return blocks


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", "replace")


def _pdf(data: bytes) -> list[tuple[int | None, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise DocumentError("PDF support needs `pip install pypdf`") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("PDF is password-protected.")
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    except DocumentError:
        raise
    except Exception as e:  # noqa: BLE001
        raise DocumentError(f"Could not read PDF: {e}") from e


def _docx(data: bytes) -> list[tuple[int | None, str]]:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise DocumentError("DOCX support needs `pip install python-docx`") from e
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise DocumentError(f"Could not read DOCX: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:                      # tables often hold the personal data
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [(None, "\n\n".join(x for x in parts if x.strip()))]


_SENT_RE = re.compile(r"(?<=[.!?။])\s+(?=[A-Z0-9\u1000-\u109F(\"'])|\n+")


def sentences(text: str) -> list[str]:
    """Cheap sentence splitter (Latin + Burmese '။' terminator)."""
    return [s.strip() for s in _SENT_RE.split(text) if len(s.strip()) > 2]
